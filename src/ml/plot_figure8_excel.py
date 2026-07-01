import openpyxl
from openpyxl.chart import BarChart, Reference
from openpyxl.chart.series import DataPoint
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter

wb = openpyxl.Workbook()
ws = wb.active
ws.title = "ML Results"

# Headers
headers = ['Model', 'Accuracy', 'Precision', 'Recall', 'F1 Score']
for col, h in enumerate(headers, 1):
    cell = ws.cell(row=1, column=col, value=h)
    cell.font = Font(bold=True)
    cell.alignment = Alignment(horizontal='center')

# Data
data = [
    ['Log',  0.929, 0.960, 0.959, 0.959],
    ['XGB',  0.900, 0.938, 0.948, 0.942],
    ['LGBM', 0.900, 0.938, 0.948, 0.942],
    ['CB',   0.906, 0.948, 0.944, 0.945],
    ['RF',   0.903, 0.941, 0.948, 0.944],
    ['DT',   0.887, 0.945, 0.926, 0.934],
    ['SVM',  0.479, 0.925, 0.441, 0.588],
    ['KNN',  0.871, 0.947, 0.904, 0.924],
    ['NB',   0.900, 0.920, 0.970, 0.944],
]

for row, d in enumerate(data, 2):
    for col, val in enumerate(d, 1):
        cell = ws.cell(row=row, column=col, value=val)
        if col > 1:
            cell.number_format = '0.000'
            cell.alignment = Alignment(horizontal='center')

# Column widths
ws.column_dimensions['A'].width = 8
for col in 'BCDE':
    ws.column_dimensions[col].width = 12

# Bar chart
chart = BarChart()
chart.type = "col"
chart.grouping = "clustered"
chart.title = "Performance Comparison of ML Models on Clinical Lung Cancer Dataset"
chart.y_axis.title = ""
chart.x_axis.title = ""
chart.style = 10
chart.width = 22
chart.height = 13
chart.y_axis.scaling.min = 0.3
chart.y_axis.scaling.max = 1.05

# Add 4 series
series_labels = ['Accuracy', 'Precision', 'Recall', 'F1 Score']
colors = ['4472C4', 'ED7D31', '70AD47', 'FFC000']

for i, (label, color) in enumerate(zip(series_labels, colors), 2):
    data_ref = Reference(ws, min_col=i, min_row=1, max_row=10)
    chart.add_data(data_ref, titles_from_data=True)
    chart.series[-1].graphicalProperties.solidFill = color
    chart.series[-1].graphicalProperties.line.solidFill = color

# X-axis labels
cats = Reference(ws, min_col=1, min_row=2, max_row=10)
chart.set_categories(cats)

ws.add_chart(chart, "G1")

wb.save('Figure8_SMOTE.xlsx')
print("Saved: Figure8_SMOTE.xlsx")
