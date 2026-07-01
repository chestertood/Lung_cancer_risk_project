"""Final grammar fix — current file has 2 novelty paras (p8,p9); correct indices confirmed."""
import sys; sys.stdout.reconfigure(encoding='utf-8')
from docx import Document
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

FILE = 'M-ID267600_v2_RiskScore.docx'

doc = Document(FILE)
P = doc.paragraphs

# verify state
assert 'To address these gaps' in P[8].text and 'To address these gaps' in P[9].text, 'Unexpected structure'
assert 'While multi-view analysis' in P[36].text, f'p36={P[36].text[:50]}'
assert 'patients data' in P[54].text, f'p54={P[54].text[:50]}'
assert 'falls short' in P[82].text, f'p82={P[82].text[:50]}'
assert 'most optimal' in P[86].text, f'p86={P[86].text[:50]}'
assert 'two branch' in P[104].text, f'p104={P[104].text[:50]}'
assert 'dection' in P[110].text, f'p110={P[110].text[:50]}'
print('Assertions passed.')

# ── helpers ───────────────────────────────────────────────────────────────
def red(r_el):
    rPr = r_el.find(qn('w:rPr'))
    if rPr is None:
        rPr = OxmlElement('w:rPr'); r_el.insert(0, rPr)
    for c in rPr.findall(qn('w:color')): rPr.remove(c)
    col = OxmlElement('w:color'); col.set(qn('w:val'), 'FF0000'); rPr.append(col)

def settext(r_el, text):
    for t in r_el.findall(qn('w:t')): r_el.remove(t)
    t = OxmlElement('w:t'); t.set(qn('xml:space'), 'preserve'); t.text = text; r_el.append(t)

def fix_substr(para_idx, old, new, label=''):
    for r in P[para_idx].runs:
        if old in r.text:
            red(r._r); settext(r._r, r.text.replace(old, new, 1))
            print(f'OK   {label or para_idx}: {repr(old)[:40]} -> {repr(new)[:40]}')
            return True
    print(f'FAIL {label or para_idx}: {repr(old)[:50]}')
    return False

# ── 1. Delete duplicate novelty (p9) ─────────────────────────────────────
P[9]._p.getparent().remove(P[9]._p)
P = doc.paragraphs  # refresh
print('Deleted duplicate novelty. p9 now:', P[9].text[:60])

# ── 2. [36→35] multi-view: Axiel→Axial, emphasizes→emphasises, remove will ──
# after deletion, indices shift -1 for everything after p9
# Current p36 → becomes p35
idx_mv = 35
assert 'While multi-view analysis' in P[idx_mv].text
fix_substr(idx_mv, 'Axiel', 'Axial', '[35] Axial')
fix_substr(idx_mv, 'emphasizes', 'emphasises', '[35] emphasises')
# "this study will " split across runs — find run ending with "will" or containing " will "
for r in P[idx_mv].runs:
    if ' will ' in r.text:
        fix_substr(idx_mv, ' will ', ' ', '[35] remove will (inline)')
        break
    elif r.text.strip() == 'will':
        red(r._r); settext(r._r, '')
        print('OK   [35] remove will (standalone run)')
        break

# ── 3. [54→53] ML branch: patients → patient ────────────────────────────
fix_substr(53, 'patients', 'patient', '[53] patient')

# ── 4. [82→81] NB: falls short → achieves, compared to → than ───────────
fix_substr(81, 'it falls short of', 'it achieves', '[81] achieves')
fix_substr(81, 'compared to the ensemble', 'than the ensemble', '[81] than')

# ── 5. [86→85] CB: most optimal → optimal ────────────────────────────────
# "making them most" is in one run or split; try both
if not fix_substr(85, 'making them most\xa0', 'making them\xa0', '[85] remove most (xa0)'):
    fix_substr(85, 'making them most ', 'making them ', '[85] remove most (space)')

# ── 6. [104→103] conclusion: two branch → two-branch ────────────────────
fix_substr(103, 'two branch', 'two-branch', '[103] two-branch')

# ── 7. [110→109] clinical sig: dection → detection ───────────────────────
fix_substr(109, 'dection', 'detection', '[109] detection')

doc.save(FILE)
print(f'\nSaved {FILE}')

# ── verify ────────────────────────────────────────────────────────────────
d = Document(FILE)
checks = [
    (8,  'dual-pipeline framework', True),
    (9,  'This study aims', True),          # objectives back to p9
    (35, 'Axial View', True),
    (35, 'Axiel', False),
    (35, 'emphasises', True),
    (53, 'structured patient data', True),
    (53, 'patients data', False),
    (81, 'it achieves', True),
    (85, 'making them\xa0optimal', True),
    (103,'two-branch parallel pipeline', True),
    (109,'early detection', True),
    (109,'dection', False),
]
print()
ok_all = True
for i, s, expect in checks:
    found = s in d.paragraphs[i].text
    ok = found == expect
    ok_all = ok_all and ok
    status = 'OK ' if ok else 'XX '
    print(f'  {status} [{i}] {"FOUND" if found else "ABSENT"} {s!r}')
print()
print('ALL OK' if ok_all else 'SOME FAILED')
