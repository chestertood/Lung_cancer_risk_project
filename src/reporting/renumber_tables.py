import sys
from docx import Document
from docx.oxml.ns import qn
sys.stdout.reconfigure(encoding='utf-8')

FILE = 'M-ID267600_v2_RiskScore.docx'
doc = Document(FILE)
P = doc.paragraphs

# ---------- 1. move risk caption + table back to Methods (after the 'illustrates' paragraph) ----------
cap = None; destD = None
for p in P:
    t = p.text.strip()
    if t.startswith('Table 8 Representative'):
        cap = p._p
    if 'illustrates representative cases' in t:
        destD = p._p
assert cap is not None and destD is not None, (cap, destD)
tbl = cap.getnext()
assert tbl.tag == qn('w:tbl'), tbl.tag
cap.getparent().remove(cap)
tbl.getparent().remove(tbl)
# insert after destD: caption then table
destD.addnext(tbl)
destD.addnext(cap)
doc.save(FILE)

# ---------- 2. renumber (plain, keep formatting) ----------
def replace_plain(par, old, new):
    for run in par.runs:
        if old in run.text:
            run.text = run.text.replace(old, new, 1)
            return True
    return False

doc = Document(FILE)
P = doc.paragraphs

# (locator substring, old, new)
edits = [
    ('comparison across all models is presented in', 'Table 2.',           'Table 3.'),
    ('illustrates representative cases',             'Table 8 illustrates', 'Table 2 illustrates'),
    ('present a comprehensive summary of the machine','Table 2 and',        'Table 3 and'),
    ('Performance comparison of various machine learning models','Table 2 Performance','Table 3 Performance'),
    ('Further investigation of the CB model',        'Table 3-5',           'Tables 4-6'),
    ('Varying Iteration',                            'Table3 Evaluation',   'Table 4 Evaluation'),
    ('Varying Eta',                                  'Table4 Evaluation',   'Table 5 Evaluation'),
    ('Varying Depth',                                'Table5 Evaluation',   'Table 6 Evaluation'),
    ('experimental results summarized in',           'Tables 6 and 7',      'Tables 7 and 8'),
    ('experimental results summarized in',           'see Table 6',         'see Table 7'),
    ('Performance comparison of various deep learning models','Table 6 Performance','Table 7 Performance'),
    ("Performance evaluation of YOLOv8n on the",     'Table7 Performance',  'Table 8 Performance'),
    ('under doctor-verified annotation',             '(Table 7)',           '(Table 8)'),
    ('Representative examples of the proposed integrated','Table 8 Representative','Table 2 Representative'),
]

for loc, old, new in edits:
    done = False
    for p in P:
        if loc in p.text and old in p.text:
            if replace_plain(p, old, new):
                done = True
                print(f'OK   {old!r} -> {new!r}')
                break
    if not done:
        print(f'FAIL {old!r}  (loc {loc[:30]!r})')

doc.save(FILE)
print('\nSaved.')

# ---------- verify ----------
d2 = Document(FILE)
import re
print('\n--- Table mentions after renumber ---')
for i, p in enumerate(d2.paragraphs):
    for m in re.finditer(r'Tables?\s*\d+(\s*(and|-|,|to)\s*\d+)?', p.text):
        s = max(0, m.start()-35)
        print(f'  ...{p.text[s:m.end()+3]}...')
