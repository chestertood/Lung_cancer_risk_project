# -*- coding: utf-8 -*-
"""
After user inserted new 'Table 4 AUC comparison', renumber colliding tables
(CB hyperparam 4-6 -> 5-7; DL 7,8 -> 8,9) and add a human-like reference to
the new AUC table at the end of the ML results paragraph.
"""
import sys
sys.stdout.reconfigure(encoding='utf-8')
from docx import Document
from docx.shared import RGBColor
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

FILE = 'M-ID267600_v2_RiskScore.docx'
doc = Document(FILE)
P = doc.paragraphs

def settext(r_el, text):
    for t in r_el.findall(qn('w:t')): r_el.remove(t)
    t = OxmlElement('w:t'); t.set(qn('xml:space'), 'preserve'); t.text = text; r_el.append(t)

def repl_run(para_idx, run_idx, old, new, label):
    """Exact-phrase replace inside a run, preserving its existing color."""
    r = P[para_idx].runs[run_idx]
    if old not in r.text:
        print(f'FAIL {label}: {old!r} not in p[{para_idx}]r[{run_idx}] = {r.text[:50]!r}')
        return False
    settext(r._r, r.text.replace(old, new, 1))
    print(f'OK   {label}: {old!r} -> {new!r}')
    return True

# ── Renumber (exact phrases, no cascade) ──────────────────────────────────
repl_run(86, 1, 'Tables 4-6',      'Tables 5-7',      '[86] CB tables range')
repl_run(89, 0, 'Table 4',         'Table 5',         '[89] iter caption')
repl_run(90, 0, 'Table 5',         'Table 6',         '[90] eta caption')
repl_run(91, 0, 'Table 6',         'Table 7',         '[91] depth caption')
repl_run(94, 1, 'Tables 7 and 8',  'Tables 8 and 9',  '[94] DL range')
repl_run(94, 3, 'Table 7)',        'Table 8)',        '[94] see Table')
repl_run(96, 0, 'Table 7',         'Table 8',         '[96] DL perf caption')
repl_run(97, 1, 'ble 8 ',          'ble 9 ',          '[97] YOLOv8n caption')
repl_run(100, 1, '(Table 8)',      '(Table 9)',       '[100] doctor-verified ref')

# ── Add human-like AUC reference to end of para 79 (ML results) ──────────
AUC_SENT = (
    ' The discriminative ability of the models, reported as the area under the ROC '
    'curve in Table 4, reinforces this pattern: the Log model again ranked highest '
    'with a mean AUC of 0.937, followed closely by the ensemble methods RF (0.922) '
    'and CB (0.915), all comfortably exceeding the 0.85 threshold adopted in this '
    'study. In contrast, the SVM and DT models recorded the lowest AUC values of '
    '0.678 and 0.795 respectively, indicating weaker separability between malignant '
    'and benign cases.'
)
p79 = P[79]
run = p79.add_run(AUC_SENT)
run.font.color.rgb = RGBColor(0xFF, 0x00, 0x00)
print('\nAdded AUC reference sentence to end of para [79].')

doc.save(FILE)
print(f'Saved {FILE}')

# ── verify ────────────────────────────────────────────────────────────────
d = Document(FILE)
print('\n--- verify table refs ---')
import re
for i in [86, 89, 90, 91, 94, 96, 97, 100]:
    t = d.paragraphs[i].text
    refs = re.findall(r'Tables?\s*\d+(?:\s*(?:and|-|to)\s*\d+)?', t)
    print(f'[{i}] {refs}')
print('\nAUC sentence present:', 'mean AUC of 0.937' in d.paragraphs[79].text)
