import sys, copy
from docx import Document
from docx.text.paragraph import Paragraph
from docx.shared import RGBColor, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
import docx_edit as de

sys.stdout.reconfigure(encoding='utf-8')
FILE = 'M-ID267600_v2_RiskScore.docx'
RED = RGBColor(0xFF, 0x00, 0x00)

doc = Document(FILE)
P = doc.paragraphs

anchor   = P[68]._p     # Figure 7 caption
head_t   = P[55]._p     # heading template ('Model evaluation and deployment')
body_t   = P[56]._p     # body template
parent   = P[68]._parent

# ---------- text ----------
HEAD = 'Integrated Risk Scoring'

B = ('To translate the complementary outputs of the two pipelines into a single, '
     'clinically interpretable measure, this study proposes an integrated risk score. '
     'Although the object detection and machine learning pipelines are trained independently '
     'on domain-appropriate datasets, both can be applied to the same patient at inference time: '
     'the object detection model contributes the predicted probability of malignancy, P(malignant), '
     'from the detected nodule, while the machine learning model contributes the predicted cancer '
     "probability, P(cancer), from the patient's clinical and symptom data. The two probabilities are "
     'weighted equally and combined into a composite score on a 0-100 scale, as defined in Equation 11:')

EQ = 'Risk Score = (P(malignant) x 50) + (P(cancer) x 50)          (11)'

D = ('Here P(malignant) and P(cancer) each range from 0 to 1, yielding a final score between 0 and 100, '
     'where higher values indicate greater estimated risk. The equal 50:50 weighting reflects the absence '
     'of prior evidence favouring either modality and can be recalibrated once validation data become '
     'available. This formulation allows the framework to remain informative even when one modality is '
     'uncertain; for example, when no nodule is detected (P(malignant) = 0), the clinical branch still '
     'contributes a risk estimate, mirroring evidence that clinical predictors retain value in the absence '
     'of visible nodules. The use of complementary imaging and clinical predictors for risk stratification '
     'is consistent with established nodule-malignancy models [34, 35] and recent multimodal fusion '
     'approaches [36]. It should be emphasised that, because the two branches were trained on separate '
     'datasets, the integrated risk score is presented here as a proposed framework; its empirical '
     'validation requires a multimodal cohort containing CT scans, clinical records, and confirmed '
     'diagnoses for the same patients, as noted in the limitations. Table 8 illustrates representative '
     'cases showing how the score behaves across differing imaging and clinical inputs.')

CAP = 'Table 8 Representative examples of the proposed integrated risk score.'

def new_para(tmpl, text, center=False):
    el = copy.deepcopy(tmpl)
    p = Paragraph(el, parent)
    de.set_para_red(p, text)
    if center:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    return el, p

# insert paragraphs in reverse -> final order A,B,C,D,E
elA,_ = new_para(head_t, HEAD)
elB,_ = new_para(body_t, B)
elC,pC = new_para(body_t, EQ, center=True)
elD,_ = new_para(body_t, D)
elE,pE = new_para(body_t, CAP)

anchor.addnext(elE)
anchor.addnext(elD)
anchor.addnext(elC)
anchor.addnext(elB)
anchor.addnext(elA)

# ---------- table ----------
rows = [
    ['Case', 'P(malignant)\n(Imaging)', 'P(cancer)\n(Clinical)', 'Risk Score\n(0-100)', 'Risk Level'],
    ['A', '0.00', '0.70', '35.0', 'Moderate'],
    ['B', '0.85', '0.60', '72.5', 'High'],
    ['C', '0.20', '0.15', '17.5', 'Low'],
]
tbl = doc.add_table(rows=len(rows), cols=5)
tbl.style = 'Table Grid'
for r, rowvals in enumerate(rows):
    for c, val in enumerate(rowvals):
        cell = tbl.cell(r, c)
        cell.text = ''
        para = cell.paragraphs[0]
        run = para.add_run(val)
        run.font.color.rgb = RED
        run.font.size = Pt(10)
        if r == 0:
            run.font.bold = True
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
# move table right after caption paragraph E
pE._p.addnext(tbl._tbl)

# ---------- references [35],[36] ----------
last_ref = doc.paragraphs[138]._p   # Mayo [34]
ref35 = ('McWilliams A, Tammemagi MC, Mayo JR, Roberts H, Liu G, Soghrati K, Yasufuku K, Martel S, '
         'Laberge F, Gingras M, Atkar-Khattra S. Probability of cancer in pulmonary nodules detected on '
         'first screening CT. New England Journal of Medicine. 2013 Sep 5;369(10):910-919.')
ref36 = ('Huang H, Zheng D, Chen H, Wang Y, Chen C, Xu L, Li G, Wang Y, He X, Li W. Fusion of CT images '
         'and clinical variables based on deep learning for predicting invasiveness risk of stage I lung '
         'adenocarcinoma. Medical Physics. 2022 Oct;49(10):6384-6394.')
for txt in (ref36, ref35):  # reverse so order ends 35 then 36
    el = copy.deepcopy(last_ref)
    p = Paragraph(el, doc.paragraphs[138]._parent)
    de.set_para_red(p, txt)
    last_ref.addnext(el)

doc.save(FILE)
print('Saved', FILE)

# verify
d2 = Document(FILE)
for i in range(68, 75):
    print(f'[{i}]', d2.paragraphs[i].text[:70])
print('tables:', len(d2.tables))
print('last refs:')
for p in d2.paragraphs[-3:]:
    print('  ', p.text[:70])
